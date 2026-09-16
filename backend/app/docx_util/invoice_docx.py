"""Invoice merge fields and .docx writers."""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from .proofing import _set_default_proofing_language_en_gb
from .quote_table import (
    INVOICE_MERGE_SLOT_COUNT,
    _INVOICE_SLOT_TOKEN_RE,
    format_gbp_pence,
)

def _invoice_line_type_label(line_type: str) -> str:
    if line_type == "fee":
        return "Fee"
    if line_type == "disbursement":
        return "Disbursement"
    if line_type == "vat":
        return "VAT"
    return line_type.replace("_", " ").title()


def invoice_line_merge_fields(
    *,
    invoice_number: str,
    invoice_date: date,
    bill_to_name: str | None,
    lines: list[dict[str, object]],
    total_pence: int,
    max_slots: int = INVOICE_MERGE_SLOT_COUNT,
) -> dict[str, str]:
    """Indexed merge codes for invoice line tables in a firm template."""
    fields: dict[str, str] = {
        "[INVOICE_NUMBER]": invoice_number,
        "[INVOICE_DATE]": invoice_date.strftime("%d %B %Y"),
        "[INVOICE_BILL_TO]": (bill_to_name or "").strip(),
        "[INVOICE_TOTAL]": format_gbp_pence(total_pence),
    }
    net_total = vat_total = 0
    used = 0
    for i, raw in enumerate(lines, start=1):
        if i > max_slots:
            break
        used = i
        tag = f"{i:02d}"
        line_type = str(raw.get("line_type") or "")
        description = str(raw.get("description") or "")
        amount_pence = int(raw.get("amount_pence") or 0)
        tax_pence = int(raw.get("tax_pence") or 0)
        gross = amount_pence + tax_pence
        net_total += amount_pence
        vat_total += tax_pence
        fields[f"[INVOICE_{tag}_TYPE]"] = _invoice_line_type_label(line_type)
        fields[f"[INVOICE_{tag}_DESCRIPTION]"] = description
        fields[f"[INVOICE_{tag}_NET]"] = format_gbp_pence(amount_pence)
        fields[f"[INVOICE_{tag}_VAT]"] = format_gbp_pence(tax_pence)
        fields[f"[INVOICE_{tag}_TOTAL]"] = format_gbp_pence(gross)
    fields["[INVOICE_NET_TOTAL]"] = format_gbp_pence(net_total)
    fields["[INVOICE_VAT_TOTAL]"] = format_gbp_pence(vat_total)
    for i in range(used + 1, max_slots + 1):
        tag = f"{i:02d}"
        fields[f"[INVOICE_{tag}_TYPE]"] = ""
        fields[f"[INVOICE_{tag}_DESCRIPTION]"] = ""
        fields[f"[INVOICE_{tag}_NET]"] = ""
        fields[f"[INVOICE_{tag}_VAT]"] = ""
        fields[f"[INVOICE_{tag}_TOTAL]"] = ""
    return fields

def strip_empty_invoice_table_rows(doc_bytes: bytes) -> bytes:
    """Remove invoice-table rows with no merged content (empty cells or leftover slot placeholders)."""
    import io

    from docx import Document

    doc = Document(io.BytesIO(doc_bytes))
    changed = False

    def _row_blank(cells: list) -> bool:
        texts = [(c.text or "").strip() for c in cells]
        if any(
            _INVOICE_SLOT_TOKEN_RE.match(t)
            for t in texts
        ):
            return True
        return not any(texts)

    def _process_table(table: Any) -> None:
        nonlocal changed
        remove_indices: list[int] = []
        for ri, row in enumerate(table.rows):
            if ri == 0:
                continue
            if _row_blank(row.cells):
                remove_indices.append(ri)
        for ri in reversed(remove_indices):
            table._tbl.remove(table.rows[ri]._tr)
            changed = True

    for table in doc.tables:
        _process_table(table)
    for section in doc.sections:
        for hf in (
            section.header,
            section.footer,
            section.even_page_header,
            section.even_page_footer,
            section.first_page_header,
            section.first_page_footer,
        ):
            if hf.is_linked_to_previous:
                continue
            for table in hf.tables:
                _process_table(table)

    if not changed:
        return doc_bytes
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

def write_invoice_docx(
    path: Path,
    *,
    firm_trading_name: str,
    firm_registered_name: str | None,
    firm_addr_line1: str | None,
    firm_addr_line2: str | None,
    firm_town_city: str | None,
    firm_county: str | None,
    firm_postcode: str | None,
    invoice_number: str,
    invoice_date: date,
    case_number: str,
    client_name: str | None,
    matter_description: str,
    fee_earner_name: str | None,
    bill_to_name: str | None,
    lines: list[dict[str, object]],
    total_pence: int,
) -> None:
    """Write a client invoice .docx from structured invoice data."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    def _fmt_pence(p: int) -> str:
        val = p / 100
        sign = "-" if val < 0 else ""
        return f"{sign}£{abs(val):,.2f}"

    def _line_type_label(line_type: str) -> str:
        if line_type == "fee":
            return "Fee"
        if line_type == "disbursement":
            return "Disbursement"
        if line_type == "vat":
            return "VAT"
        return line_type.replace("_", " ").title()

    doc = Document()
    _set_default_proofing_language_en_gb(doc)

    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    firm_line = (firm_trading_name or "").strip()
    if firm_registered_name and firm_registered_name.strip() and firm_registered_name.strip() != firm_line:
        firm_line = f"{firm_line} ({firm_registered_name.strip()})" if firm_line else firm_registered_name.strip()
    if firm_line:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(firm_line)
        r.bold = True
        r.font.size = Pt(14)

    addr_parts = [firm_addr_line1, firm_addr_line2, firm_town_city, firm_county, firm_postcode]
    addr_text = ", ".join(x.strip() for x in addr_parts if x and str(x).strip())
    if addr_text:
        addr_para = doc.add_paragraph()
        addr_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        addr_run = addr_para.add_run(addr_text)
        addr_run.font.size = Pt(10)
        addr_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run("INVOICE")
    title_run.bold = True
    title_run.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        f"{invoice_number}  ·  {invoice_date.strftime('%d %B %Y')}"
    )
    meta_run.font.size = Pt(11)

    doc.add_paragraph()

    if bill_to_name and bill_to_name.strip():
        bill_para = doc.add_paragraph()
        bill_para.add_run("Bill to: ").bold = True
        bill_para.add_run(bill_to_name.strip())

    matter_bits = [case_number]
    if client_name and client_name.strip():
        matter_bits.append(client_name.strip())
    if matter_description and matter_description.strip():
        matter_bits.append(matter_description.strip())
    matter_para = doc.add_paragraph("Matter: " + " — ".join(matter_bits))
    matter_para.runs[0].font.size = Pt(10)

    if fee_earner_name and fee_earner_name.strip():
        fe_para = doc.add_paragraph(f"Fee earner: {fee_earner_name.strip()}")
        fe_para.runs[0].font.size = Pt(10)

    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ("Type", "Description", "Net", "VAT", "Total")
    for i, label in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = label
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 1 else WD_ALIGN_PARAGRAPH.LEFT

    net_total = vat_total = 0
    for raw in lines:
        line_type = str(raw.get("line_type") or "")
        description = str(raw.get("description") or "")
        amount_pence = int(raw.get("amount_pence") or 0)
        tax_pence = int(raw.get("tax_pence") or 0)
        gross = amount_pence + tax_pence
        net_total += amount_pence
        vat_total += tax_pence
        row = table.add_row()
        row.cells[0].text = _line_type_label(line_type)
        row.cells[1].text = description
        row.cells[2].text = _fmt_pence(amount_pence)
        row.cells[3].text = _fmt_pence(tax_pence)
        row.cells[4].text = _fmt_pence(gross)
        for j in range(2, 5):
            row.cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()

    summary = doc.add_table(rows=3, cols=2)
    summary.style = "Table Grid"
    summary_rows = [
        ("Net total", _fmt_pence(net_total)),
        ("VAT total", _fmt_pence(vat_total)),
        ("Invoice total", _fmt_pence(total_pence)),
    ]
    for i, (label, value) in enumerate(summary_rows):
        summary.rows[i].cells[0].text = label
        summary.rows[i].cells[1].text = value
        if i == 2:
            summary.rows[i].cells[0].paragraphs[0].runs[0].bold = True
            summary.rows[i].cells[1].paragraphs[0].runs[0].bold = True
        summary.rows[i].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
