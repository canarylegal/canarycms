"""Quote fee table markers, xlsx grid inserts, and table presentation."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from .proofing import _set_default_proofing_language_en_gb

PRECEDENT_BODY_MARKER = "[PRECEDENT_BODY]"
QUOTE_FEE_TABLE_MARKER = "[QUOTE_FEE_TABLE]"
QUOTE_TABLE_MARKERS = (QUOTE_FEE_TABLE_MARKER, PRECEDENT_BODY_MARKER)
QUOTE_MERGE_SLOT_COUNT = 25
_QUOTE_SLOT_TOKEN_RE = re.compile(r"^\[QUOTE_\d{2}_(?:LABEL|AMOUNT|VAT)\]$", re.IGNORECASE)
INVOICE_MERGE_SLOT_COUNT = 25
_INVOICE_SLOT_TOKEN_RE = re.compile(
    r"^\[INVOICE_\d{2}_(?:TYPE|DESCRIPTION|NET|VAT|TOTAL)\]$",
    re.IGNORECASE,
)
COMPLETION_MERGE_SLOT_COUNT = 50
_COMPLETION_SLOT_TOKEN_RE = re.compile(
    r"^\[COMPLETION_\d{2}_(?:DESCRIPTION|DEBIT|CREDIT)\]$",
    re.IGNORECASE,
)


def insert_xlsx_grid_table_at_marker(doc_bytes: bytes, grid: "XlsxGrid") -> bytes:
    """Replace ``[QUOTE_FEE_TABLE]`` / ``[PRECEDENT_BODY]`` with a Word table built from an xlsx grid."""
    import io
    from copy import deepcopy

    from docx import Document
    from docx.oxml.ns import qn

    from app.xlsx_util import XlsxGrid

    if not isinstance(grid, XlsxGrid):
        raise TypeError("grid must be XlsxGrid")

    doc = Document(io.BytesIO(doc_bytes))
    body = doc.element.body
    p_tag = qn("w:p")
    t_tag = qn("w:t")
    tbl_tag = qn("w:tbl")
    sect_pr_tag = qn("w:sectPr")

    marker_para = None
    for el in list(body):
        if el.tag != p_tag:
            continue
        text = "".join((t.text or "") for t in el.iter(t_tag))
        if any(m in text for m in QUOTE_TABLE_MARKERS):
            marker_para = el
            break

    tmp = Document()
    nrows = len(grid.rows) or 1
    ncols = max((len(r) for r in grid.rows), default=1) or 1
    table = tmp.add_table(rows=nrows, cols=ncols)
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    for r_i, row in enumerate(grid.rows):
        for c_i in range(ncols):
            val = row[c_i] if c_i < len(row) else ""
            cell = table.rows[r_i].cells[c_i]
            cell.text = val
            if (r_i, c_i) in grid.bold:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True
    for r0, c0, r1, c1 in grid.merges:
        try:
            table.rows[r0].cells[c0].merge(table.rows[r1].cells[c1])
        except (IndexError, ValueError):
            pass

    tbl_el = deepcopy(table._tbl)

    if marker_para is not None:
        parent = marker_para.getparent()
        idx = list(parent).index(marker_para)
        parent.insert(idx, tbl_el)
        parent.remove(marker_para)
    else:
        sect_pr = next((el for el in body if el.tag == sect_pr_tag), None)
        if sect_pr is not None:
            idx = list(body).index(sect_pr)
            body.insert(idx, tbl_el)
        else:
            body.append(tbl_el)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def format_gbp_pence(pence: int | None) -> str:
    if pence is None:
        return ""
    negative = pence < 0
    pence = abs(pence)
    pounds = pence / 100
    text = f"£{pounds:,.2f}"
    return f"-{text}" if negative else text


def insert_quote_fee_table_at_marker(
    doc_bytes: bytes,
    rows: list[tuple[str, str | None, bool, bool]],
) -> bytes:
    """Insert a two-column fee table. Each row: (label, amount_display, is_bold, amount_right)."""
    import io
    from copy import deepcopy

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document(io.BytesIO(doc_bytes))
    body = doc.element.body
    p_tag = qn("w:p")
    t_tag = qn("w:t")
    sect_pr_tag = qn("w:sectPr")

    marker_para = None
    for el in list(body):
        if el.tag != p_tag:
            continue
        text = "".join((t.text or "") for t in el.iter(t_tag))
        if any(m in text for m in QUOTE_TABLE_MARKERS):
            marker_para = el
            break

    tmp = Document()
    nrows = max(len(rows), 1)
    table = tmp.add_table(rows=nrows, cols=2)
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    data = rows if rows else [(" ", "")]
    for r_i, (label, amount, is_bold, amount_right) in enumerate(data):
        left = table.rows[r_i].cells[0]
        right = table.rows[r_i].cells[1]
        left.text = label
        right.text = amount or ""
        for cell, right_align in ((left, False), (right, amount_right)):
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if right_align else WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    if is_bold:
                        run.bold = True

    tbl_el = deepcopy(table._tbl)
    if marker_para is not None:
        parent = marker_para.getparent()
        idx = list(parent).index(marker_para)
        parent.insert(idx, tbl_el)
        parent.remove(marker_para)
    else:
        sect_pr = next((el for el in body if el.tag == sect_pr_tag), None)
        if sect_pr is not None:
            idx = list(body).index(sect_pr)
            body.insert(idx, tbl_el)
        else:
            body.append(tbl_el)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def strip_precedent_body_marker(doc_bytes: bytes) -> bytes:
    """Remove any leftover ``[PRECEDENT_BODY]`` token from a .docx so it never renders literally.

    Replaces the token text with an empty string inside any paragraph whose combined run-text
    contains it (keeps the surrounding paragraph for layout / spacing). Used as a safety net for
    the blank-letter compose path where the splice is skipped, and as defence-in-depth after the
    splice. Handles the marker even when split across multiple ``w:t`` runs (a common artefact of
    editing in ONLYOFFICE / Word).
    """
    import io
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(io.BytesIO(doc_bytes))
    body = doc.element.body
    p_tag = qn("w:p")
    t_tag = qn("w:t")
    changed = False
    for el in list(body.iter(p_tag)):
        runs_text = "".join((t.text or "") for t in el.iter(t_tag))
        if PRECEDENT_BODY_MARKER not in runs_text:
            continue
        cleaned = runs_text.replace(PRECEDENT_BODY_MARKER, "")
        for r in list(el.findall(qn("w:r"))):
            el.remove(r)
        if cleaned:
            r = el.makeelement(qn("w:r"), {})
            t = el.makeelement(qn("w:t"), {qn("xml:space"): "preserve"})
            t.text = cleaned
            r.append(t)
            el.append(r)
        changed = True
    if not changed:
        return doc_bytes
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

_QUOTE_TABLE_HEADER_FILL = "D0DAEA"
_QUOTE_TABLE_SECTION_FILL = "EEF2F8"
_QUOTE_TABLE_TOTAL_FILL = "D0DAEA"
_QUOTE_TABLE_BORDER_COLOR = "B8C4D4"
_QUOTE_TABLE_BORDER_LIGHT = "D8DEE8"


def _docx_set_cell_shading(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _docx_set_cell_margin_dxa(
    cell: Any,
    *,
    top: int = 80,
    bottom: int = 80,
    left: int = 120,
    right: int = 120,
) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def _docx_set_cell_borders(
    cell: Any,
    *,
    top: dict[str, object] | None = None,
    bottom: dict[str, object] | None = None,
    left: dict[str, object] | None = None,
    right: dict[str, object] | None = None,
) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for side, spec in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        if spec is None:
            continue
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), str(spec.get("val", "single")))
        el.set(qn("w:sz"), str(spec.get("sz", 4)))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), str(spec.get("color", "auto")))
        tc_borders.append(el)
    tc_pr.append(tc_borders)


def _docx_set_table_width_pct(table: Any, pct: int = 5000) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    for existing in tbl_pr.findall(qn("w:tblW")):
        tbl_pr.remove(existing)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(pct))
    tbl_w.set(qn("w:type"), "pct")
    tbl_pr.append(tbl_w)


def _docx_normalize_table_width(table: Any) -> None:
    """Ensure a table has a single ``w:tblW`` (Word uses the first when duplicates exist)."""
    from copy import deepcopy

    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        return
    widths = tbl_pr.findall(qn("w:tblW"))
    if len(widths) <= 1:
        return
    chosen = None
    for width in widths:
        if width.get(qn("w:type")) == "pct":
            chosen = width
            break
    if chosen is None:
        for width in reversed(widths):
            if width.get(qn("w:w")) not in (None, "0"):
                chosen = width
                break
    if chosen is None:
        chosen = widths[-1]
    for width in widths:
        tbl_pr.remove(width)
    tbl_pr.append(deepcopy(chosen))


def _docx_set_table_layout_fixed(table: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _docx_table_grid_col_widths(table: Any) -> list[int]:
    from docx.oxml.ns import qn

    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is None:
        return []
    out: list[int] = []
    for col in grid.findall(qn("w:gridCol")):
        raw = col.get(qn("w:w"))
        if raw is not None:
            out.append(int(raw))
    return out


def _docx_set_table_grid_col_widths(table: Any, widths: list[int]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl_pr = tbl.tblPr
        insert_at = list(tbl).index(tbl_pr) + 1 if tbl_pr is not None else 0
        tbl.insert(insert_at, grid)
    for col in list(grid.findall(qn("w:gridCol"))):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)


# Default fee-table column widths (3.85", 1.35", 1.35") when the template omits ``w:tblGrid``.
_QUOTE_FEE_TABLE_GRID_DXA = (5544, 1944, 1944)


def _docx_style_table_paragraph(
    para: Any,
    *,
    bold: bool = False,
    size_pt: float = 10,
    alignment: Any | None = None,
) -> None:
    from docx.shared import Pt

    if alignment is not None:
        para.alignment = alignment
    text = para.text
    para.clear()
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)


def _docx_quote_border(*, light: bool = False, strong: bool = False) -> dict[str, object]:
    if strong:
        return {"val": "single", "sz": 8, "color": _QUOTE_TABLE_BORDER_COLOR}
    color = _QUOTE_TABLE_BORDER_LIGHT if light else _QUOTE_TABLE_BORDER_COLOR
    return {"val": "single", "sz": 4, "color": color}


def _find_quote_fee_table(doc: Any) -> Any | None:
    for table in doc.tables:
        if not table.rows:
            continue
        if len(table.rows[0].cells) < 3:
            continue
        hdr = (table.rows[0].cells[0].text or "").strip().lower()
        if hdr == "description":
            return table
    return None


def _docx_copy_table_width(from_table: Any, to_table: Any) -> None:
    """Copy ``w:tblW`` from one table to another so layout matches the template."""
    from copy import deepcopy

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    _docx_normalize_table_width(from_table)
    src_pr = from_table._tbl.tblPr
    if src_pr is None:
        return
    src_w = src_pr.find(qn("w:tblW"))
    if src_w is None:
        return
    dst_tbl = to_table._tbl
    dst_pr = dst_tbl.tblPr
    if dst_pr is None:
        dst_pr = OxmlElement("w:tblPr")
        dst_tbl.insert(0, dst_pr)
    for existing in dst_pr.findall(qn("w:tblW")):
        dst_pr.remove(existing)
    dst_pr.append(deepcopy(src_w))


def _docx_lock_quote_fee_table_layout(table: Any) -> list[int]:
    """Preserve template column widths so merged content cannot expand the fee table."""
    _docx_normalize_table_width(table)
    _docx_set_table_layout_fixed(table)
    ncols = len(table.rows[0].cells) if table.rows else 0
    grid = _docx_table_grid_col_widths(table)
    if len(grid) != ncols:
        if ncols == 3:
            grid = list(_QUOTE_FEE_TABLE_GRID_DXA)
        elif ncols == 2:
            grid = [_QUOTE_FEE_TABLE_GRID_DXA[0], sum(_QUOTE_FEE_TABLE_GRID_DXA[1:])]
        else:
            grid = []
        if grid:
            _docx_set_table_grid_col_widths(table, grid)
    return grid


def _docx_summary_grid_from_fee_grid(fee_grid: list[int], summary_cols: int) -> list[int]:
    if summary_cols <= 0 or not fee_grid:
        return []
    if summary_cols == len(fee_grid):
        return list(fee_grid)
    if summary_cols == 2 and len(fee_grid) >= 3:
        return [fee_grid[0], sum(fee_grid[1:3])]
    if summary_cols == 2 and len(fee_grid) == 2:
        return list(fee_grid)
    total = sum(fee_grid)
    if summary_cols == 1:
        return [total]
    per = max(total // summary_cols, 1)
    return [per] * (summary_cols - 1) + [total - per * (summary_cols - 1)]


def _find_quote_summary_tables(doc: Any, *, exclude: Any | None = None) -> list[Any]:
    """Tables for quote totals (e.g. grand total) below the main fee table."""
    tables: list[Any] = []
    for table in doc.tables:
        if table is exclude or not table.rows:
            continue
        ncols = len(table.rows[0].cells)
        if ncols not in (2, 3):
            continue
        blob = " ".join((cell.text or "") for row in table.rows for cell in row.cells).lower()
        if "grand total" in blob or "quote_grand_total" in blob:
            tables.append(table)
    return tables


def _sync_quote_summary_table_widths(doc: Any, fee_table: Any, *, fee_grid: list[int]) -> None:
    for summary in _find_quote_summary_tables(doc, exclude=fee_table):
        _docx_copy_table_width(fee_table, summary)
        _docx_set_table_layout_fixed(summary)
        summary_cols = len(summary.rows[0].cells) if summary.rows else 0
        grid = _docx_summary_grid_from_fee_grid(fee_grid, summary_cols)
        if grid:
            _docx_set_table_grid_col_widths(summary, grid)


def _style_quote_fee_table_header_row(table: Any, *, preserve_layout: bool = False) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    col_widths = (Inches(3.85), Inches(1.35), Inches(1.35))
    hdr = table.rows[0].cells
    labels = ("Description", "Amount", "VAT")
    for i, label in enumerate(labels):
        cell = hdr[i]
        cell.text = label
        para = cell.paragraphs[0]
        _docx_style_table_paragraph(
            para,
            bold=True,
            alignment=WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT,
        )
        _docx_set_cell_shading(cell, _QUOTE_TABLE_HEADER_FILL)
        _docx_set_cell_margin_dxa(cell)
        _docx_set_cell_borders(
            cell,
            top=_docx_quote_border(strong=True),
            bottom=_docx_quote_border(strong=True),
            left=_docx_quote_border() if i == 0 else None,
            right=_docx_quote_border() if i == 2 else None,
        )
    if not preserve_layout:
        for i, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = width


def _style_quote_fee_table_data_row(
    row: Any,
    *,
    line_kind: str = "item",
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    is_section = line_kind == "section_header"
    is_total_row = line_kind in ("subtotal", "total")

    if is_section:
        for ci, cell in enumerate(row.cells[:3]):
            if not cell.paragraphs:
                cell.add_paragraph()
            para = cell.paragraphs[0]
            if ci == 0:
                _docx_style_table_paragraph(para, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
            elif para.text.strip():
                para.clear()
            _docx_set_cell_shading(cell, _QUOTE_TABLE_SECTION_FILL)
            _docx_set_cell_margin_dxa(cell)
            _docx_set_cell_borders(
                cell,
                bottom=_docx_quote_border(light=True),
                left=_docx_quote_border() if ci == 0 else None,
                right=_docx_quote_border() if ci == 2 else None,
            )
        return

    fill = _QUOTE_TABLE_TOTAL_FILL if is_total_row else None
    for ci, cell in enumerate(row.cells[:3]):
        if not cell.paragraphs:
            cell.add_paragraph()
        para = cell.paragraphs[0]
        _docx_style_table_paragraph(
            para,
            bold=is_total_row,
            alignment=WD_ALIGN_PARAGRAPH.RIGHT if ci > 0 else WD_ALIGN_PARAGRAPH.LEFT,
        )
        _docx_set_cell_margin_dxa(cell)
        if fill:
            _docx_set_cell_shading(cell, fill)
        _docx_set_cell_borders(
            cell,
            bottom=_docx_quote_border(light=True),
            left=_docx_quote_border() if ci == 0 else None,
            right=_docx_quote_border() if ci == 2 else None,
        )


def apply_quote_table_presentation(doc_bytes: bytes, lines: list[Any]) -> bytes:
    """Apply header, section, and total styling to the merged quote fee table."""
    import io

    from docx import Document

    doc = Document(io.BytesIO(doc_bytes))
    table = _find_quote_fee_table(doc)
    if table is None:
        return doc_bytes

    fee_grid = _docx_lock_quote_fee_table_layout(table)

    _style_quote_fee_table_header_row(table, preserve_layout=True)

    for ri, line in enumerate(lines, start=1):
        if ri >= len(table.rows):
            break
        kind = line.line_kind.value if hasattr(line.line_kind, "value") else str(line.line_kind)
        _style_quote_fee_table_data_row(
            table.rows[ri],
            line_kind=kind,
        )

    _sync_quote_summary_table_widths(doc, table, fee_grid=fee_grid)
    _docx_lock_quote_fee_table_layout(table)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def strip_empty_quote_table_rows(doc_bytes: bytes) -> bytes:
    """Remove fee-table rows with no merged content (empty cells or leftover slot placeholders)."""
    import io

    from docx import Document

    doc = Document(io.BytesIO(doc_bytes))
    changed = False

    def _row_blank(cells: list) -> bool:
        if len(cells) < 2:
            return all(not (c.text or "").strip() for c in cells)
        label = (cells[0].text or "").strip()
        amount = (cells[1].text or "").strip()
        vat = (cells[2].text or "").strip() if len(cells) > 2 else ""
        if (
            _QUOTE_SLOT_TOKEN_RE.match(label)
            or _QUOTE_SLOT_TOKEN_RE.match(amount)
            or _QUOTE_SLOT_TOKEN_RE.match(vat)
        ):
            return True
        return not label and not amount and not vat

    def _process_table(table: Any) -> None:
        nonlocal changed
        remove_indices: list[int] = []
        for ri, row in enumerate(table.rows):
            if ri == 0:
                continue
            if _row_blank(row.cells):
                remove_indices.append(ri)
        for ri in reversed(remove_indices):
            table._tbl.remove(table.rows[ri]._tr)
            changed = True

    for table in doc.tables:
        _process_table(table)
    for section in doc.sections:
        for hf in (
            section.header,
            section.footer,
            section.even_page_header,
            section.even_page_footer,
            section.first_page_header,
            section.first_page_footer,
        ):
            if hf.is_linked_to_previous:
                continue
            for table in hf.tables:
                _process_table(table)

    if not changed:
        return doc_bytes
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

def write_quote_template_docx(path: Path, *, slots: int = QUOTE_MERGE_SLOT_COUNT) -> None:
    """Write a minimal quote .docx with indexed fee-table merge slots (no letter precedent)."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    _set_default_proofing_language_en_gb(doc)

    doc.add_paragraph("[ORG_AND_ADDRESS_BLOCK]")
    doc.add_paragraph("")
    doc.add_paragraph("[DATE]")
    doc.add_paragraph("")
    doc.add_paragraph("Re: [MATTER_DESCRIPTION] — [CASE_REF]")
    doc.add_paragraph("")
    doc.add_paragraph("Dear [CONTACT_LETTER_DEAR]")
    doc.add_paragraph("")
    doc.add_paragraph(
        "Thank you for instructing us. Set out below is our estimate of costs based on a property value of "
        "[QUOTE_PROPERTY_VALUE]."
    )
    doc.add_paragraph("")

    table = doc.add_table(rows=1 + slots, cols=3)
    _docx_set_table_width_pct(table)
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    _style_quote_fee_table_header_row(table)

    for i in range(1, slots + 1):
        tag = f"{i:02d}"
        row = table.rows[i].cells
        row[0].text = f"[QUOTE_{tag}_LABEL]"
        row[1].paragraphs[0].text = f"[QUOTE_{tag}_AMOUNT]"
        row[2].paragraphs[0].text = f"[QUOTE_{tag}_VAT]"
        _style_quote_fee_table_data_row(table.rows[i])

    doc.add_paragraph("")
    closing = doc.add_paragraph("Yours faithfully")
    closing.runs[0].font.size = Pt(11)
    doc.add_paragraph("")
    doc.add_paragraph("[FEE_EARNER]")
    doc.add_paragraph("[FIRM_TRADING_NAME]")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def write_quote_template_docx_bytes(*, slots: int = QUOTE_MERGE_SLOT_COUNT) -> bytes:
    import tempfile

    fd, tmp_name = tempfile.mkstemp(suffix=".docx")
    tmp = Path(tmp_name)
    try:
        import os

        os.close(fd)
        write_quote_template_docx(tmp, slots=slots)
        return tmp.read_bytes()
    finally:
        tmp.unlink(missing_ok=True)
