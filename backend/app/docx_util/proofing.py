"""Docx proofing language, finalize helpers, and blank/email writers."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

_W_MAIN_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# Match self-closing or opening ``w:lang`` / ``w:themeFontLang`` elements only.
_LANG_ELEM_RE = re.compile(
    r"<(?:(?:w):)?(?:lang|themeFontLang)\b[^>]*/?>|<(?:(?:w):)?(?:lang|themeFontLang)\b[^>]*>",
    re.IGNORECASE,
)
_LANG_ATTR_RE = re.compile(
    r'(\s(?:(?:w):)?(?:val|eastAsia|bidi)=["\'])([^"\']*)(["\'])',
    re.IGNORECASE,
)
_STYLES_OPEN_RE = re.compile(r"(<(?:(?:w):)?styles\b[^>]*>)", re.IGNORECASE)
_SETTINGS_OPEN_RE = re.compile(r"(<(?:(?:w):)?settings\b[^>]*>)", re.IGNORECASE)
_DOC_DEFAULTS_BLOCK_RE = re.compile(
    r"<(?:(?:w):)?docDefaults\b[^>]*>.*?</(?:(?:w):)?docDefaults>",
    re.IGNORECASE | re.DOTALL,
)
_DOC_DEFAULTS_BLOCK = (
    "<w:docDefaults><w:rPrDefault><w:rPr>"
    '<w:lang w:val="en-GB" w:eastAsia="en-GB" w:bidi="en-GB"/>'
    "</w:rPr></w:rPrDefault></w:docDefaults>"
)
_THEME_FONT_LANG_BLOCK = '<w:themeFontLang w:val="en-GB" w:eastAsia="en-GB" w:bidi="en-GB"/>'


def _coerce_lang_attr_to_en_gb(val: str | None) -> str:
    v = (val or "").strip()
    if not v:
        return "en-GB"
    low = v.replace("_", "-").lower()
    if low in ("en-us", "en", "en-us-x"):
        return "en-GB"
    if low == "en-gb":
        return "en-GB"
    return v


def _patch_lang_element_xml(elem_xml: str) -> str:
    def attr_repl(match: re.Match[str]) -> str:
        return match.group(1) + _coerce_lang_attr_to_en_gb(match.group(2)) + match.group(3)

    return _LANG_ATTR_RE.sub(attr_repl, elem_xml)


def _patch_ooxml_lang_text(text: str, *, part_filename: str = "") -> str:
    """Patch language tags in-place without rewriting the whole OOXML tree.

    ElementTree ``tostring`` on large ``styles.xml`` parts rewrites namespaces and can
    bloat or corrupt complex firm letterheads — ONLYOFFICE then fails to open the file.
    """
    text = _LANG_ELEM_RE.sub(lambda m: _patch_lang_element_xml(m.group(0)), text)
    base = part_filename.rsplit("/", 1)[-1]
    if base == "styles.xml" and not re.search(r"<(?:(?:w):)?docDefaults\b", text, re.IGNORECASE):
        m = _STYLES_OPEN_RE.search(text)
        if m:
            text = text[: m.end()] + _DOC_DEFAULTS_BLOCK + text[m.end() :]
    if base == "settings.xml" and not re.search(r"<(?:(?:w):)?themeFontLang\b", text, re.IGNORECASE):
        m = _SETTINGS_OPEN_RE.search(text)
        if m:
            text = text[: m.end()] + _THEME_FONT_LANG_BLOCK + text[m.end() :]
    return text


def _patch_ooxml_lang_bytes(raw: bytes, *, part_filename: str = "") -> bytes:
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw
    patched = _patch_ooxml_lang_text(text, part_filename=part_filename)
    if patched == text:
        return raw
    return patched.encode("utf-8")


def ensure_docx_proofing_language_en_gb_bytes(src_bytes: bytes) -> bytes:
    """Force British English as the document language across Word OOXML parts.

    ONLYOFFICE and Word read ``w:docDefaults``, ``w:themeFontLang``, and per-run ``w:lang``.
    Without a full pass, DS often shows “English (United States)” and rewrites saves as en-US.
    """
    import io
    import zipfile

    if not src_bytes.startswith(b"PK"):
        return src_bytes
    try:
        zin = zipfile.ZipFile(io.BytesIO(src_bytes), "r")
    except zipfile.BadZipFile:
        return src_bytes

    patches: dict[str, bytes] = {}
    for info in zin.infolist():
        if not info.filename.startswith("word/") or not info.filename.endswith(".xml"):
            continue
        raw = zin.read(info.filename)
        patched = _patch_ooxml_lang_bytes(raw, part_filename=info.filename)
        if patched != raw:
            patches[info.filename] = patched

    if not patches:
        zin.close()
        return src_bytes

    out_buf = io.BytesIO()
    with zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = patches.get(info.filename, zin.read(info.filename))
            zout.writestr(info, data)
    zin.close()
    return out_buf.getvalue()


def normalize_onlyoffice_persisted_docx_bytes(
    data: bytes,
    *,
    filename: str | None = None,
    mime_type: str | None = None,
) -> bytes:
    """Re-save a .docx exported via ONLYOFFICE ``downloadAs`` so it can be reopened reliably.

    ``downloadAs`` can leave orphan relationship parts and OOXML that triggers ONLYOFFICE
    ``changesError`` on the next open. Round-tripping through python-docx strips those artefacts
    while preserving body content, tables, headers, and embedded media.
    """
    import io

    name = (filename or "").lower()
    mt = (mime_type or "").split(";", 1)[0].strip().lower()
    if not (name.endswith(".docx") or mt == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"):
        return data
    if not data.startswith(b"PK"):
        return data
    try:
        from docx import Document

        doc = Document(io.BytesIO(data))
        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()
    except Exception:
        return data


def finalize_stored_docx_bytes(
    data: bytes,
    *,
    filename: str | None = None,
    mime_type: str | None = None,
) -> bytes:
    """Apply en-GB language normalisation before persisting a .docx from ONLYOFFICE / WebDAV."""
    name = (filename or "").lower()
    mt = (mime_type or "").split(";", 1)[0].strip().lower()
    if name.endswith(".docx") or mt == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return ensure_docx_proofing_language_en_gb_bytes(data)
    return data


def _set_default_proofing_language_en_gb(doc: Any) -> None:
    """Set OOXML default run language to en-GB for new documents.

    ONLYOFFICE/Word use ``w:docDefaults`` for default document / proofing language. python-docx often
    omits ``docDefaults`` until we create it.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    styles_el = doc.styles.element
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is None:
        dd = OxmlElement("w:docDefaults")
        styles_el.insert(0, dd)
    rpd = dd.find(qn("w:rPrDefault"))
    if rpd is None:
        rpd = OxmlElement("w:rPrDefault")
        dd.insert(0, rpd)
    rpr = rpd.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpd.append(rpr)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "en-GB")
    lang.set(qn("w:eastAsia"), "en-GB")
    lang.set(qn("w:bidi"), "en-GB")


def write_blank_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    _set_default_proofing_language_en_gb(doc)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def write_blank_email_precedent_docx(path: Path) -> None:
    """Plain-text-friendly default e-mail body (reserved ``BLANK_EMAIL`` precedent)."""
    from docx import Document

    doc = Document()
    _set_default_proofing_language_en_gb(doc)

    lines = [
        "[CONTACT_LETTER_DEAR]",
        "",
        "[PRECEDENT_BODY]",
        "",
        "[CONTACT_LETTER_SIGN_OFF]",
        "",
        "[FEE_EARNER]",
        "[FEE_EARNER_JOB_TITLE]",
        "[FIRM_TRADING_NAME]",
        "",
        "---",
        "Our ref: [FEE_EARNER_INITIALS]/[CASE_REF]    Your ref: [CONTACT_REF]",
        "[DATE]",
    ]
    for line in lines:
        doc.add_paragraph(line)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def write_quote_email_precedent_docx(path: Path) -> None:
    """Plain-text-friendly e-mail body for sending a quote (Thunderbird / mailto / Graph)."""
    from docx import Document

    doc = Document()
    _set_default_proofing_language_en_gb(doc)

    lines = [
        "[CONTACT_LETTER_DEAR]",
        "",
        "Thank you for your enquiry. Please find attached our quote for [MATTER_DESCRIPTION].",
        "",
        (
            "The quote sets out the work we propose to undertake and our fees. It also shows VAT and "
            "any disbursements where applicable. If you would like to proceed, or if you have any "
            "questions, please reply to this e-mail."
        ),
        "",
        "We look forward to hearing from you.",
        "",
        "",
        "Kind regards",
        "",
        "[FEE_EARNER]",
        "[FEE_EARNER_JOB_TITLE]",
        "[FIRM_TRADING_NAME]",
        "",
        "---",
        "Our ref: [FEE_EARNER_INITIALS]/[CASE_REF]    Your ref: [CONTACT_REF]",
        "[DATE]",
    ]
    for line in lines:
        doc.add_paragraph(line)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))

def extract_plain_text_from_docx_bytes(data: bytes) -> str:
    """Best-effort plain text from a .docx for e-mail body (M365 Graph)."""
    from io import BytesIO

    from docx import Document

    doc = Document(BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    t = (p.text or "").strip()
                    if t:
                        parts.append(t)
    return "\n\n".join(parts) if parts else ""


def write_client_account_reconcile_report_docx(
    path: Path,
    *,
    firm_trading_name: str,
    firm_registered_name: str | None,
    client_bank_account_name: str | None,
    client_bank_sort_code: str | None,
    client_bank_account_number_last4: str | None,
    client_bank_account_number: str | None = None,
    period_end_date: date,
    ledger_client_total_pence: int,
    ledger_office_total_pence: int,
    bank_statement_balance_pence: int,
    difference_pence: int,
    prepared_by_name: str | None,
    prepared_at: datetime | None,
    approved_by_name: str | None,
    approved_at: datetime | None,
    notes: str | None,
    status: str,
) -> None:
    """Write a client account reconcile report .docx for month-end sign-off."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    def _fmt_pence(p: int) -> str:
        val = p / 100
        sign = "-" if val < 0 else ""
        return f"{sign}£{abs(val):,.2f}"

    def _fmt_dt(dt: datetime | None) -> str:
        if dt is None:
            return "—"
        local = dt
        if local.tzinfo is not None:
            local = local.replace(tzinfo=None)
        return local.strftime("%d %B %Y %H:%M")

    doc = Document()
    _set_default_proofing_language_en_gb(doc)

    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("CLIENT ACCOUNT RECONCILE REPORT")
    run.bold = True
    run.font.size = Pt(16)

    firm_line = (firm_trading_name or "").strip()
    if firm_registered_name and firm_registered_name.strip() and firm_registered_name.strip() != firm_line:
        firm_line = f"{firm_line} ({firm_registered_name.strip()})" if firm_line else firm_registered_name.strip()
    if firm_line:
        firm_para = doc.add_paragraph()
        firm_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        firm_run = firm_para.add_run(firm_line)
        firm_run.font.size = Pt(12)

    period_para = doc.add_paragraph()
    period_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    period_run = period_para.add_run(f"Period ended: {period_end_date.strftime('%d %B %Y')}")
    period_run.font.size = Pt(11)

    doc.add_paragraph()

    bank_bits: list[str] = []
    if client_bank_account_name:
        bank_bits.append(client_bank_account_name.strip())
    if client_bank_sort_code:
        bank_bits.append(f"Sort code {client_bank_sort_code.strip()}")
    acct = (client_bank_account_number or "").strip()
    if acct:
        bank_bits.append(f"Account {acct}")
    elif client_bank_account_number_last4:
        bank_bits.append(f"Account •••• {client_bank_account_number_last4.strip()}")
    if bank_bits:
        bank_para = doc.add_paragraph("Client bank account: " + " · ".join(bank_bits))
        bank_para.runs[0].font.size = Pt(10)

    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    rows_data = [
        ("Ledger client total (all matters)", _fmt_pence(ledger_client_total_pence)),
        ("Bank statement closing balance", _fmt_pence(bank_statement_balance_pence)),
        ("Difference (bank minus ledger)", _fmt_pence(difference_pence)),
        ("Office ledger total (reference)", _fmt_pence(ledger_office_total_pence)),
        ("Status", status.capitalize()),
    ]
    for i, (label, value) in enumerate(rows_data):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value

    doc.add_paragraph()

    prep_para = doc.add_paragraph(f"Prepared by: {prepared_by_name or '—'}")
    prep_para.runs[0].font.size = Pt(10)
    prep_at = doc.add_paragraph(f"Prepared at: {_fmt_dt(prepared_at)}")
    prep_at.runs[0].font.size = Pt(10)

    appr_para = doc.add_paragraph(f"Approved by: {approved_by_name or '—'}")
    appr_para.runs[0].font.size = Pt(10)
    appr_at = doc.add_paragraph(f"Approved at: {_fmt_dt(approved_at)}")
    appr_at.runs[0].font.size = Pt(10)

    if notes and notes.strip():
        doc.add_paragraph()
        notes_heading = doc.add_paragraph("Notes")
        notes_heading.runs[0].bold = True
        for line in notes.strip().splitlines():
            doc.add_paragraph(line)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
